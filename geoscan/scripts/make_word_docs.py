"""Build professionally formatted Word versions of the review documents.

    python scripts/make_word_docs.py <package_dir>

1. Creates a styled pandoc reference document (fonts, heading hierarchy, bordered tables with a
   shaded header row, A4 margins, footer with page numbers).
2. Pre-processes the Markdown: title block from metadata, formulas converted to native Word
   equations (LaTeX -> OMML via pandoc), table of contents on its own page.
3. Runs pandoc, then Microsoft Word (COM) to refresh the table of contents, fit tables to the
   page width, repeat table header rows, and export a PDF preview.
"""

from __future__ import annotations

import copy
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PANDOC = shutil.which("pandoc") or r"C:\Program Files\Pandoc\pandoc.exe"
NAVY = RGBColor(0x1F, 0x38, 0x64)
FONT = "Calibri"

# inline-code formula -> LaTeX (rendered as Word equations)
FORMULAS = {
    "`C(λ) = R_L + (λ − λ_L)·(R_R − R_L)/(λ_R − λ_L)`":
        r"$$C(\lambda)=R_L+(\lambda-\lambda_L)\,\frac{R_R-R_L}{\lambda_R-\lambda_L}$$",
    "`CR(λ) = R(λ) / C(λ)`": r"$$CR(\lambda)=\frac{R(\lambda)}{C(\lambda)}$$",
    "`y(x) = y₀ + d₁(x − x₀) + a(x − x₀)(x − x₁)`": r"$y(x)=y_0+d_1(x-x_0)+a(x-x_0)(x-x_1)$",
    "`d₁ = (y₁ − y₀)/(x₁ − x₀)`": r"$d_1=\frac{y_1-y_0}{x_1-x_0}$",
    "`a = (d₂ − d₁)/(x₂ − x₀)`": r"$a=\frac{d_2-d_1}{x_2-x_0}$",
    "`λ_min = (x₀ + x₁)/2 − d₁/(2a)`": r"$\lambda_{min}=\frac{x_0+x_1}{2}-\frac{d_1}{2a}$",
    "`BD = 1 − CR(λ_min)`": r"$BD=1-CR(\lambda_{min})$",
    "`BD̄ = Σ wᵢ BDᵢ / Σ wᵢ`": r"$\overline{BD}=\frac{\sum_i w_i\,BD_i}{\sum_i w_i}$",
    "`wᵢ = 1/σᵢ²`": r"$w_i=1/\sigma_i^2$",
    "`σ̄ = (Σ wᵢ)^(−1/2)`": r"$\bar{\sigma}=\left(\sum_i w_i\right)^{-1/2}$",
    "`SNR = BD̄ / σ̄`": r"$SNR=\overline{BD}/\bar{\sigma}$",
    "`F_core = G_core * (w·BD) / G_core * w`":
        r"$F_{core}=\frac{G_{core}*(w\cdot BD)}{G_{core}*w}$",
    "`W_halo = G_halo * (G_core * w)`": r"$W_{halo}=G_{halo}*(G_{core}*w)$",
    "`conf = 1 − exp(−W_halo / θ)`": r"$conf=1-\exp(-W_{halo}/\theta)$",
    "`F = (G_core * F_core) · conf`": r"$F=(G_{core}*F_{core})\cdot conf$",
    "`a = clip(F / P98(F > 0 inside the block), 0, 1)`":
        r"$a=\mathrm{clip}\left(\frac{F}{P_{98}(F>0\ \mathrm{inside\ block})},\,0,\,1\right)$",
    "`QI = e₁₁² / (e₁₀ · e₁₂)`": r"$QI=\frac{e_{11}^2}{e_{10}\,e_{12}}$",
    "`SI = e₁₃ / e₁₂`": r"$SI=\frac{e_{13}}{e_{12}}$",
    "`QSI = N(0.55·N(QI) + 0.45·N(SI))`": r"$QSI=N\left(0.55\,N(QI)+0.45\,N(SI)\right)$",
    "`N(X) = clip((X − P5)/(P98 − P5), 0, 1)`": r"$N(X)=\mathrm{clip}\left(\frac{X-P_5}{P_{98}-P_5},0,1\right)$",
    "`a_qtz = clip((QSI − P15)/(P98 − P15), 0, 1)`":
        r"$a_{qtz}=\mathrm{clip}\left(\frac{QSI-P_{15}}{P_{98}-P_{15}},0,1\right)$",
    "`CI = e₁₃ / e₁₄`": r"$CI=e_{13}/e_{14}$",
    "`MI = e₁₂ · e₁₄³ / e₁₃⁴`": r"$MI=\frac{e_{12}\,e_{14}^3}{e_{13}^4}$",
    "`H = cos Z cos S + sin Z sin S cos(A − φ)`": r"$H=\cos Z\,\cos S+\sin Z\,\sin S\,\cos(A-\varphi)$",
    "`sym = z_centre − (z_left + z_right)/2`": r"$sym=z_{centre}-\frac{z_{left}+z_{right}}{2}$",
    "`step = |z_left − z_right|/2`": r"$step=\frac{\lvert z_{left}-z_{right}\rvert}{2}$",
    "step > |sym|": r"step > $\lvert sym\rvert$",
    "`S_prox = exp(−d / 250 m)`": r"$S_{prox}=\exp(-d/250\ \mathrm{m})$",
}
CPI_BLOCK = ("```\nH   = 0.30·a_musc + 0.28·a_qtz + 0.22·a_goss + 0.20·a_chl\n"
             "CPI = 1_block · (0.70·H + 0.30·S_prox)\n```")
CPI_LATEX = (r"$$H=0.30\,a_{musc}+0.28\,a_{qtz}+0.22\,a_{goss}+0.20\,a_{chl}$$" "\n\n"
             r"$$CPI=\mathbf{1}_{block}\cdot\left(0.70\,H+0.30\,S_{prox}\right)$$")
PAGE_BREAK = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'


def _set_font(style, size=None, bold=None, color=None, italic=None):
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), FONT)
    for att in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rfonts.attrib.pop(qn(att), None)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    if color is not None:
        style.font.color.rgb = color


def _table_style(doc):
    """Bordered 'Table' style (pandoc's default table style) with a shaded header row."""
    st = doc.styles["Table"]
    el = st.element
    tblpr = el.find(qn("w:tblPr"))
    if tblpr is None:
        tblpr = OxmlElement("w:tblPr")
        el.append(tblpr)
    for old in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), "A6A6A6")
        borders.append(b)
    tblpr.append(borders)
    mar = OxmlElement("w:tblCellMar")
    for side, w in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        mar.append(m)
    tblpr.append(mar)
    for old in el.findall(qn("w:tblStylePr")):
        el.remove(old)
    first = OxmlElement("w:tblStylePr")
    first.set(qn("w:type"), "firstRow")
    rpr = OxmlElement("w:rPr")
    b = OxmlElement("w:b")
    rpr.append(b)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), "1F3864")
    rpr.append(col)
    first.append(rpr)
    tcpr = OxmlElement("w:tcPr")
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "D9E2F3")
    tcpr.append(shd)
    first.append(tcpr)
    el.append(first)
    # table text 9.5 pt
    trpr = el.get_or_add_rPr() if hasattr(el, "get_or_add_rPr") else None
    st.font.size = Pt(9.5)


def _footer(doc, text):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.2))
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
    p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"{text}  ·  page ")
    r.font.size, r.font.color.rgb, r.font.name = Pt(8), RGBColor(0x59, 0x59, 0x59), FONT
    for code in ("PAGE",):
        for kind, val in (("begin", None), ("instr", f" {code} "), ("end", None)):
            run = p.add_run()
            run.font.size = Pt(8)
            if kind == "instr":
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = val
                run._r.append(it)
            else:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                run._r.append(fc)


def reference_doc(path: Path, footer_text: str):
    raw = subprocess.run([PANDOC, "-o", str(path), "--print-default-data-file", "reference.docx"], check=True)
    doc = Document(path)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in doc.styles:
            _set_font(doc.styles[name], size=10.5, color=RGBColor(0x1A, 0x1A, 0x1A))
            pf = doc.styles[name].paragraph_format
            pf.space_after, pf.line_spacing = Pt(4), 1.1
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5), ("Heading 4", 10.5)):
        _set_font(doc.styles[name], size=size, bold=True, color=NAVY, italic=False)
        doc.styles[name].paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        doc.styles[name].paragraph_format.space_after = Pt(4)
    _set_font(doc.styles["Title"], size=22, bold=True, color=NAVY)
    for name in ("Subtitle", "Date", "Author"):
        if name in doc.styles:
            _set_font(doc.styles[name], size=12, color=RGBColor(0x40, 0x40, 0x40), italic=False)
    if "TOC Heading" in doc.styles:
        _set_font(doc.styles["TOC Heading"], size=14, bold=True, color=NAVY)
    for name in ("Source Code", "Verbatim Char"):
        if name in doc.styles:
            doc.styles[name].font.name = "Consolas"
            doc.styles[name].font.size = Pt(9)
    _table_style(doc)
    _footer(doc, footer_text)
    doc.save(path)


LIST = re.compile(r"^\s*([-*]|\d+\.)\s")
LABEL = re.compile(r"^\*\*[^*]+:\*\*")


def separate_blocks(s: str) -> str:
    """Pandoc needs a blank line before a list and between consecutive '**Label:**' lines;
    without it the list (or the labels) are merged into the preceding paragraph."""
    out, prev = [], ""
    for line in s.split("\n"):
        starts_list = bool(LIST.match(line)) and not line.startswith("  ")
        if prev.strip() and not LIST.match(prev) and not prev.startswith("|") and not prev.startswith("  ") and (
                starts_list or (LABEL.match(line) and LABEL.match(prev))):
            out.append("")
        out.append(line)
        prev = line
    return "\n".join(out)


def prepare_markdown(src: Path, title: str, subtitle: str, date: str, toc_break: bool) -> str:
    s = src.read_text(encoding="utf-8")
    s = re.sub(r"\A# .*\n", "", s)                       # title comes from metadata
    # front-matter lines of the methodology belong on the title page, not in the body
    s = re.sub(r"^\*\*Blocks:\*\* .*\n\*\*Software:\*\* .*\n", "", s, flags=re.M)
    s = separate_blocks(s)
    for k, v in FORMULAS.items():
        s = s.replace(k, v)
    s = s.replace(CPI_BLOCK, CPI_LATEX)
    s = re.sub(r"^---\s*$", "", s, flags=re.M)            # horizontal rules -> nothing (headings separate sections)
    left = [k for k in FORMULAS if k in s]
    assert not left, f"formulas not converted: {left}"
    meta = f'---\ntitle: "{title}"\nsubtitle: "{subtitle}"\ndate: "{date}"\n---\n\n'
    body = s.lstrip()
    if toc_break:
        body = PAGE_BREAK + "\n" + body
    return meta + body


def word_postprocess(docx_path: Path, pdf_path: Path):
    """Refresh TOC, fit tables to window, repeat header rows, export PDF (Microsoft Word via COM)."""
    ps = f"""
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$d = $w.Documents.Open('{docx_path}')
if ($d.CompatibilityMode -lt 15) {{ $d.Convert() }}   # upgrade to current Word format (no "Compatibility Mode")
foreach ($t in $d.Tables) {{ $t.AutoFitBehavior(2); $t.Rows.AllowBreakAcrossPages = 0; $t.Rows.Item(1).HeadingFormat = -1 }}
foreach ($toc in $d.TablesOfContents) {{ $toc.Update() }}
$d.Fields.Update() | Out-Null
$d.Save()
$d.ExportAsFixedFormat('{pdf_path}', 17)
$d.Close()
$w.Quit()
"""
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


def build(md: Path, out_docx: Path, title: str, subtitle: str, date: str, toc: bool, preview_dir: Path):
    with tempfile.TemporaryDirectory() as td:
        td = Path(td)
        ref = td / "reference.docx"
        reference_doc(ref, title.split(" — ")[0])
        tmp = td / "input.md"
        tmp.write_text(prepare_markdown(md, title, subtitle, date, toc), encoding="utf-8")
        cmd = [PANDOC, str(tmp), "-o", str(out_docx), "--reference-doc", str(ref)]
        if toc:
            cmd += ["--toc", "--toc-depth=2", "-V", "toc-title=Contents"]
        subprocess.run(cmd, check=True)
    preview_dir.mkdir(parents=True, exist_ok=True)
    word_postprocess(out_docx.resolve(), (preview_dir / (out_docx.stem + ".pdf")).resolve())


def main():
    pkg = Path(sys.argv[1]).resolve()
    review = Path(__file__).resolve().parents[1] / "review"
    preview = pkg.parent / "_preview"
    build(review / "Phase1_Methodology.md", pkg / "01_Methodology" / "Phase1_Methodology.docx",
          "Phase 1 Methodology — Satellite Alteration Mapping, Structural Lineaments and Drill-Target Ranking",
          "Kehalpur (E) Block (Betul, MP; 102.7 km²) and Thanewasna (Chandrapur, MH; 2.26 km²) · "
          "software: geoscan v0.1.0 · for geological review",
          "September 2026", toc=True, preview_dir=preview)
    build(review / "00_READ_ME_FIRST.md", pkg / "00_READ_ME_FIRST.docx",
          "Phase 1 Review Package — Read Me First",
          "Satellite alteration mapping and drill-target ranking · Kehalpur and Thanewasna",
          "September 2026", toc=False, preview_dir=preview)
    print("done; previews in", preview)


if __name__ == "__main__":
    main()
